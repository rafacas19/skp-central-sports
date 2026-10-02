"""Score sheets: saving, the mirrored observation, the headline rating, merges,
and the dashboard pages that drive them."""

from datetime import datetime, timezone

import httpx
import pytest
import pytest_asyncio

from tests.conftest import make_user

from scouting_bot.config import settings
from scouting_bot.dashboard import auth
from scouting_bot.evaluations import MergeConflict, carry_scores, save_evaluation
from scouting_bot.models import SOURCE_WEB, Evaluation, Observation, Prospect, Session
from scouting_bot.profiles import get_profile

PASSWORD = "prueba-scouting"

# The client's example report (feedback/Perfiles_Scout__EJ.pdf), Lateral profile,
# PARTIDO 1 — the scores exactly as printed, in code order.
CLIENT_EXAMPLE = {
    "1.1": 4, "1.2": 3, "1.3": 3, "1.4": 4,
    "2.1": 4, "2.2": 3, "2.3": 4, "2.4": 4, "2.5": 4, "2.6": 4, "2.7": 3, "2.8": 3, "2.9": 4,
    "3.1": 2, "3.2": 3, "3.3": 3, "3.4": 4, "3.5": 3, "3.6": 4,
    "4.1": 4, "4.2": 2, "4.3": 4,
    "5.1": 4, "5.2": 3,
}


def _when(day: int) -> datetime:
    return datetime(2026, 9, day, 18, 0, tzinfo=timezone.utc)


async def _match(day: int = 20, home: str = "Junior", away: str = "Nacional") -> Session:
    return await Session.create(
        agent_chat_id=1, home_team=home, away_team=away, state="ended", match_date=_when(day)
    )


async def _player(name: str = "Camilo Restrepo", team: str = "Junior", **kw) -> Prospect:
    from scouting_bot.taxonomy import normalize_identity, normalize_name

    return await Prospect.create(
        agent_chat_id=1, name=name, normalized_name=normalize_identity(name),
        team=team, normalized_team=normalize_name(team), **kw,
    )


# ── Saving a sheet ───────────────────────────────────────────────────────
async def test_client_example_gives_their_numbers(storage):
    match, player = await _match(), await _player(position="Lateral izquierdo")
    ev = await save_evaluation(match, player, profile=get_profile("lateral"), raw_scores=CLIENT_EXAMPLE)

    assert ev.rating == 3.46  # 83 / 24
    obs = await Observation.get(id=ev.observation_id)
    assert (obs.source, obs.rating, obs.session_id, obs.prospect_id) == (SOURCE_WEB, 3.46, match.id, player.id)
    assert obs.side == "home" and obs.team == "Junior"
    player = await Prospect.get(id=player.id)
    assert player.latest_rating == 3.46
    assert player.decision_status == "Interesante"


async def test_saving_twice_keeps_one_sheet_and_one_observation(storage):
    match, player = await _match(), await _player()
    lateral = get_profile("lateral")
    await save_evaluation(match, player, profile=lateral, raw_scores={"1.1": "4"})
    await save_evaluation(match, player, profile=lateral, raw_scores={"1.1": "2", "1.2": "2"}, note="Bien arriba")

    assert await Evaluation.all().count() == 1
    obs = await Observation.filter(prospect_id=player.id)
    assert len(obs) == 1
    assert (obs[0].rating, obs[0].raw_quote) == (2.0, "Bien arriba")


async def test_an_empty_sheet_writes_no_rating(storage):
    match, player = await _match(), await _player()
    ev = await save_evaluation(match, player, profile=get_profile("extremo"), raw_scores={})
    assert ev.rating is None
    assert await Observation.all().count() == 0
    assert (await Prospect.get(id=player.id)).latest_rating is None


async def test_clearing_a_sheet_removes_its_observation(storage):
    match, player = await _match(), await _player()
    extremo = get_profile("extremo")
    await save_evaluation(match, player, profile=extremo, raw_scores={"1.1": 3})
    assert await Observation.all().count() == 1
    await save_evaluation(match, player, profile=extremo, raw_scores={})
    assert await Observation.all().count() == 0


async def test_a_partial_sheet_averages_only_what_was_scored(storage):
    match, player = await _match(), await _player()
    ev = await save_evaluation(
        match, player, profile=get_profile("lateral"), raw_scores={"1.1": 5, "2.1": 4, "9.9": 1}
    )
    assert ev.scores == {"1.1": 5, "2.1": 4}
    assert ev.rating == 4.5


async def test_scoring_an_older_match_does_not_overwrite_a_newer_rating(storage):
    player = await _player()
    old, new = await _match(day=10), await _match(day=20)
    lateral = get_profile("lateral")

    await save_evaluation(new, player, profile=lateral, raw_scores={"1.1": 5})
    await save_evaluation(old, player, profile=lateral, raw_scores={"1.1": 1})
    player = await Prospect.get(id=player.id)
    assert player.latest_rating == 5.0
    assert player.decision_status == "A firmar"

    # …but re-scoring the newest match does move it.
    await save_evaluation(new, player, profile=lateral, raw_scores={"1.1": 2})
    assert (await Prospect.get(id=player.id)).latest_rating == 2.0


async def test_a_goalkeeper_gets_a_single_rating(storage):
    match, keeper = await _match(), await _player("Kevin Mier", position="Portero")
    ev = await save_evaluation(match, keeper, profile=None, raw_scores={"1.1": 5}, single_rating=4.0)
    assert (ev.profile, ev.scores, ev.rating) == (None, {}, 4.0)
    assert (await Observation.get(id=ev.observation_id)).raw_quote == "Valoración del partido"


def test_switching_profile_keeps_only_criteria_of_the_same_name():
    lateral, extremo = get_profile("lateral"), get_profile("extremo")
    # Lateral 2.1 is Posicionamiento Defensivo; Extremo 2.1 is Presión — no carry.
    carried = carry_scores(lateral, extremo, {"1.1": 4, "2.1": 5, "3.1": 2})
    # 1.1 Control ≠ 1.1 Controles; 3.1 "1v1 Ofensivo" exists in both (Extremo 3.1).
    assert carried == {"3.1": 2}
    assert carry_scores(None, extremo, {"1.1": 4}) == {}


# ── Merges ───────────────────────────────────────────────────────────────
async def test_merge_moves_sheets_to_the_survivor(storage):
    keep, drop = await _player("Camilo Restrepo"), await _player("Camilo Restrepo", team="Junior FC")
    m1, m2 = await _match(day=10), await _match(day=20)
    lateral = get_profile("lateral")
    await save_evaluation(m1, keep, profile=lateral, raw_scores={"1.1": 4})
    await save_evaluation(m2, drop, profile=lateral, raw_scores={"1.1": 3})

    await storage.merge_prospects(keep.id, drop.id)
    assert await Evaluation.filter(prospect_id=keep.id).count() == 2
    assert await Observation.filter(prospect_id=keep.id).count() == 2
    assert not await Prospect.filter(id=drop.id).exists()


async def test_merge_refuses_two_sheets_for_the_same_match(storage):
    keep, drop = await _player("Camilo Restrepo"), await _player("Camilo Restrepo", team="Junior FC")
    match = await _match()
    lateral = get_profile("lateral")
    await save_evaluation(match, keep, profile=lateral, raw_scores={"1.1": 4})
    await save_evaluation(match, drop, profile=lateral, raw_scores={"1.1": 3})

    with pytest.raises(MergeConflict):
        await storage.merge_prospects(keep.id, drop.id)
    assert await Prospect.filter(id=drop.id).exists()
    assert await Evaluation.all().count() == 2
    assert await Observation.all().count() == 2


# ── Dashboard pages ──────────────────────────────────────────────────────
def _set(field: str, value) -> object:
    old = getattr(settings, field)
    object.__setattr__(settings, field, value)
    return old


@pytest_asyncio.fixture
async def client(storage):
    from scouting_bot.app import app

    old = {f: _set(f, v) for f, v in
           [("dashboard_password", PASSWORD), ("dashboard_secret", ""), ("use_mock_ai", True)]}
    auth._attempts.clear()
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as c:
        await make_user(password=PASSWORD)
        await c.post("/dashboard/login", data={"username": "tester", "password": PASSWORD})
        yield c
    for f, v in old.items():
        _set(f, v)
    auth._attempts.clear()


def _csrf(html: str) -> str:
    marker = f'name="{auth.CSRF_FIELD}" value="'
    start = html.index(marker) + len(marker)
    return html[start : html.index('"', start)]


def _url(match: Session, player: Prospect) -> str:
    return f"/dashboard/partidos/{match.id}/jugadores/{player.id}/evaluar"


async def test_sheet_route_reproduces_the_clients_report_numbers(client):
    match, player = await _match(), await _player(position="Lateral izquierdo")
    page = await client.get(_url(match, player))
    assert page.status_code == 200
    assert "Posicionamiento Defensivo" in page.text
    form = {auth.CSRF_FIELD: _csrf(page.text), "perfil": "lateral", "contextura": "si", "estatura": "si"}
    form.update({f"s-{code}": str(v) for code, v in CLIENT_EXAMPLE.items()})

    resp = await client.post(_url(match, player), data=form)
    assert resp.status_code == 303
    page = await client.get(resp.headers["location"])
    for avg in ("3,5/5", "3,7/5", "3,2/5", "3,3/5"):
        assert avg in page.text
    assert "3.46 / 5" in page.text
    assert "Evaluación guardada" in page.text

    player = await Prospect.get(id=player.id)
    assert (player.latest_rating, player.decision_status) == (3.46, "Interesante")
    ev = await Evaluation.get(prospect_id=player.id)
    assert (ev.build_ok, ev.height_ok) == (True, True)


async def test_sheet_saves_without_javascript_and_twice_is_one_sheet(client):
    match, player = await _match(), await _player(position="Extremo derecho")
    page = await client.get(_url(match, player))
    assert "Controles" in page.text  # Extremo profile picked from the position
    token = _csrf(page.text)
    for score in ("3", "5"):
        await client.post(_url(match, player), data={auth.CSRF_FIELD: token, "perfil": "extremo", "s-1.1": score})
    assert await Evaluation.all().count() == 1
    assert (await Evaluation.first()).scores == {"1.1": 5}
    assert await Observation.all().count() == 1
    page = await client.get(_url(match, player))
    assert 'value="5" checked' in page.text
    assert "incompleto" in page.text


async def test_goalkeeper_sheet_uses_the_arquero_profile(client):
    match, keeper = await _match(), await _player("Kevin Mier", position="Portero")
    page = await client.get(_url(match, keeper))
    assert "Blocaje" in page.text and "Valoración del partido" not in page.text
    # 4.1 has no description yet: a plain label, no empty disclosure.
    assert '<span class="criterion-name criterion-plain"><span class="code">4.1</span>' in page.text
    await client.post(_url(match, keeper), data={
        auth.CSRF_FIELD: _csrf(page.text), "perfil": "arquero", "s-1.5": "5", "s-2.4": "4"})
    ev = await Evaluation.first()
    assert (ev.profile, ev.rating) == ("arquero", 4.5)


async def test_a_goalkeeper_can_still_get_a_single_rating(client):
    match, keeper = await _match(), await _player("Kevin Mier", position="Portero")
    page = await client.get(_url(match, keeper) + "?perfil=ninguno")
    assert "Valoración del partido" in page.text
    await client.post(_url(match, keeper), data={auth.CSRF_FIELD: _csrf(page.text), "perfil": "", "valoracion": "4"})
    assert (await Evaluation.first()).rating == 4.0


async def test_profile_switch_warns_before_losing_scores(client):
    match, player = await _match(), await _player(position="Lateral izquierdo")
    page = await client.get(_url(match, player))
    await client.post(_url(match, player), data={auth.CSRF_FIELD: _csrf(page.text), "perfil": "lateral", "s-2.1": "5"})
    page = await client.get(_url(match, player) + "?perfil=extremo")
    assert "Perfil cambiado" in page.text
    assert "se perderán 1 puntuaciones" in page.text
    assert (await Evaluation.first()).profile == "lateral"  # nothing saved by looking


async def test_match_page_lists_players_and_opens_a_sheet_for_a_new_one(client):
    match, player = await _match(), await _player()
    await Observation.create(session=match, prospect=player, raw_quote="Buen partido", team="Junior")
    page = await client.get(f"/dashboard/partidos/{match.id}")
    assert "Evaluaciones" in page.text and _url(match, player) in page.text

    resp = await client.post(
        f"/dashboard/partidos/{match.id}/evaluar",
        data={auth.CSRF_FIELD: _csrf(page.text), "nombre": "Camilo Restrepo", "lado": "home"},
    )
    assert resp.status_code == 303
    assert resp.headers["location"].endswith(_url(match, player))  # reused, not duplicated
    assert await Prospect.all().count() == 1

    resp = await client.post(
        f"/dashboard/partidos/{match.id}/evaluar",
        data={auth.CSRF_FIELD: _csrf(page.text), "dorsal": "7", "lado": "away", "posicion": "Extremo izquierdo"},
    )
    temp = await Prospect.get(is_temporary=True)
    assert (temp.team, temp.shirt_number, temp.position) == ("Nacional", 7, "Extremo izquierdo")
    page = await client.get(resp.headers["location"])
    assert "Sin identificar (dorsal 7)" in page.text

    resp = await client.post(
        f"/dashboard/partidos/{match.id}/evaluar", data={auth.CSRF_FIELD: _csrf(page.text), "lado": "home"}
    )
    assert "aviso=faltan_datos" in resp.headers["location"]


async def test_evaluated_player_shows_his_rating_on_the_match_page(client):
    match, player = await _match(), await _player()
    await save_evaluation(match, player, profile=get_profile("lateral"), raw_scores={"1.1": 4, "1.2": 3})
    page = await client.get(f"/dashboard/partidos/{match.id}")
    assert "Ver evaluación" in page.text and "3.5 / 5" in page.text


async def test_deleting_a_sheet_removes_its_observation(client):
    match, player = await _match(), await _player()
    await save_evaluation(match, player, profile=get_profile("lateral"), raw_scores={"1.1": 4})
    page = await client.get(_url(match, player))
    resp = await client.post(_url(match, player) + "/borrar", data={"csrf": _csrf(page.text)})
    assert resp.status_code == 303
    assert await Evaluation.all().count() == 0
    assert await Observation.all().count() == 0


async def test_panel_merge_shows_the_conflict(client):
    keep, drop = await _player("Camilo Restrepo"), await _player("Camilo Restrepo", team="Junior FC")
    match = await _match()
    lateral = get_profile("lateral")
    await save_evaluation(match, keep, profile=lateral, raw_scores={"1.1": 4})
    await save_evaluation(match, drop, profile=lateral, raw_scores={"1.1": 3})
    page = await client.get(f"/dashboard/jugadores/{keep.id}/fusionar?con={drop.id}")
    resp = await client.post(
        f"/dashboard/jugadores/{keep.id}/fusionar", data={"con": drop.id, "csrf": _csrf(page.text)}
    )
    assert resp.status_code == 409
    assert "No se pudo fusionar" in resp.text
    assert await Prospect.filter(id=drop.id).exists()


async def test_sheet_post_requires_csrf(client):
    match, player = await _match(), await _player()
    resp = await client.post(_url(match, player), data={auth.CSRF_FIELD: "0" * 20, "s-1.1": "4"})
    assert resp.status_code == 400
    assert await Evaluation.all().count() == 0


async def test_sheet_has_a_tab_per_section_with_live_counts(client):
    match, player = await _match(), await _player(position="Lateral izquierdo")
    await save_evaluation(match, player, profile=get_profile("lateral"), raw_scores=CLIENT_EXAMPLE)
    page = await client.get(_url(match, player))
    counts = [page.text.split(f'data-tab="sec-{n}"')[1].split("data-count>")[1].split("<")[0]
              for n in range(1, 6)]
    assert counts == ["4/4", "9/9", "6/6", "3/3", "2/2"]
    assert 'data-tab="sec-fisico"' in page.text
    for label in ("Técnica", "Defensa", "Ataque", "Condición", "Mental", "Físico"):
        assert f'<span class="tab-label">{label}</span>' in page.text
    # Without JavaScript every section is on the page and Guardar is there.
    assert page.text.count("data-section") == 6 and "data-save>Guardar" in page.text


# ── Scores as the client's workbook (build list #1) ──────────────────────
async def test_scores_excel_follows_the_clients_workbook(client):
    import io

    import openpyxl

    from scouting_bot.models import MatchPlayer

    match = await _match()
    lateral_a = await _player("Camilo Restrepo", position="Lateral izquierdo")
    lateral_b = await _player("Luis Arias", position="Lateral derecho")
    keeper = await _player("Kevin Mier", position="Portero")
    await MatchPlayer.create(session=match, prospect=lateral_a, side="home", slot=1, shirt_number=3)
    await save_evaluation(match, lateral_a, profile=get_profile("lateral"), raw_scores=CLIENT_EXAMPLE,
                          build_ok=True, height_ok=False)
    await save_evaluation(match, lateral_b, profile=get_profile("lateral"), raw_scores={"1.1": 2})
    await save_evaluation(match, keeper, profile=None, raw_scores={}, single_rating=4.0)

    resp = await client.get(f"/dashboard/partidos/{match.id}/evaluaciones.xlsx")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Lateral", "Sin perfil"]
    ws = wb["Lateral"]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0][0] == "DETECCIÓN DE TALENTO" and rows[1][0] == "PERFIL POR POSICIÓN: LATERAL"
    assert rows[3][:5] == ["Código", "Criterio", "Descripción", "3 · Camilo Restrepo", "Luis Arias"]
    assert rows[4][0] == "1. TÉCNICA"
    control = next(r for r in rows if r[0] == "1.1")
    assert (control[1], control[3], control[4]) == ("Control", 4, 2)
    assert next(r for r in rows if r[1] == "Atlético")[3:5] == ["Sí", None]
    assert next(r for r in rows if r[1] == "Mediano")[3] == "No"
    assert next(r for r in rows if r[1] == "VALORACIÓN DEL PARTIDO (media)")[3:5] == [3.46, 2.0]
    keeper_rows = [[c.value for c in r] for r in wb["Sin perfil"].iter_rows()]
    assert ["Kevin Mier", 4.0] in keeper_rows

    page = await client.get(f"/dashboard/partidos/{match.id}")
    assert f"/dashboard/partidos/{match.id}/evaluaciones.xlsx" in page.text


async def test_scores_excel_of_an_unscored_match_still_opens(client):
    import io

    import openpyxl

    match = await _match()
    resp = await client.get(f"/dashboard/partidos/{match.id}/evaluaciones.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Evaluaciones"]


# ── AI summary reads the scores (build list #5) ──────────────────────────
async def test_the_ai_summary_gets_the_scores_and_refreshes_on_rescoring(client, monkeypatch):
    from scouting_bot.ai.mock import MockAIProvider
    from scouting_bot.dashboard import summaries

    seen = []
    original = MockAIProvider.summarize_player

    async def spy(self, observations, evaluations=None):
        seen.append((observations, evaluations))
        return await original(self, observations, evaluations)

    monkeypatch.setattr(MockAIProvider, "summarize_player", spy)
    match, player = await _match(), await _player(position="Lateral izquierdo")
    await save_evaluation(match, player, profile=get_profile("lateral"), raw_scores=CLIENT_EXAMPLE)

    page = await client.get(f"/dashboard/jugadores/{player.id}")
    assert "Puntuado por perfil en 1 partido(s)" in page.text
    observations, sheets = seen[-1]
    assert observations == []  # the sheet's stand-in text is not sent as a note
    assert sheets[0]["profile"] == "Lateral" and sheets[0]["rating"] == 3.46
    tecnica = sheets[0]["sections"][0]
    assert tecnica["section"] == "Técnica" and tecnica["average"] == "3,5"
    assert "Control" in tecnica["strengths"] and tecnica["weaknesses"] == []

    # Re-scoring changes no observation count, but the summary still refreshes.
    calls = len(seen)
    await save_evaluation(match, player, profile=get_profile("lateral"), raw_scores={"1.1": 1})
    await client.get(f"/dashboard/jugadores/{player.id}")  # stale → background refresh
    assert len(seen) == calls + 1
    await client.get(f"/dashboard/jugadores/{player.id}")  # fresh now → no new call
    assert len(seen) == calls + 1


async def test_bot_player_report_includes_the_scores(service, storage):
    match, player = await _match(), await _player(position="Lateral izquierdo")
    await save_evaluation(match, player, profile=get_profile("lateral"), raw_scores=CLIENT_EXAMPLE)
    result = await service.player_report(1, "Camilo Restrepo")
    assert "Puntuado por perfil" in result[2]
